import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_cover_certificate_abstract(output_path="CRISS_Cover_Certificate_Abstract.docx"):
    doc = docx.Document()
    
    # Configure default style to Times New Roman
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    
    # ---------------------------------------------------------
    # SECTION 1: COVER PAGE
    # ---------------------------------------------------------
    sec1 = doc.sections[0]
    sec1.top_margin = Pt(45)
    sec1.bottom_margin = Pt(14)
    sec1.left_margin = Pt(54)
    sec1.right_margin = Pt(54)
    sec1.header_distance = Pt(36)
    sec1.footer_distance = Pt(36)
    
    def add_p(text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=0, line_spacing=1.0):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    # Cover Page Content
    add_p("A MINI PROJECT REPORT ON", size=12, bold=True, space_before=0, space_after=8)
    add_p("CAMPUS RESOURCE INTELLIGENCE & SHARING SYSTEM", size=16, bold=True, space_before=6, space_after=12)
    add_p("SUBMITTED TO THE SAVITRIBAI PHULE PUNE UNIVERSITY, PUNE IN THE PARTIAL FULFILLMENT OF THE REQUIREMENTS FOR THE AWARD OF THE DEGREE OF", 
          size=11, bold=False, space_before=4, space_after=6, line_spacing=1.15)
    add_p("BACHELOR OF TECHNOLOGY", size=14, bold=True, space_before=2, space_after=2)
    add_p("IN", size=11, bold=False, space_before=0, space_after=2)
    add_p("INFORMATION TECHNOLOGY", size=15, bold=True, space_before=2, space_after=12)
    
    add_p("SUBMITTED BY", size=12, bold=False, space_before=4, space_after=6)
    add_p("[Student Name]          Exam No: [Exam Number]", size=12, bold=True, space_before=2, space_after=10)
    
    # College Logo
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(4)
    p_logo.paragraph_format.space_after = Pt(8)
    run_logo = p_logo.add_run()
    run_logo.add_picture("college_logo.jpeg", width=Inches(1.25))
    
    add_p("DEPARTMENT OF INFORMATION TECHNOLOGY", size=15, bold=True, space_before=4, space_after=4)
    add_p("JSPM’s RAJARSHI SHAHU COLLEGE OF ENGINEERING", size=12, bold=True, space_before=2, space_after=2)
    add_p("TATHAWADE, PUNE 411033", size=12, bold=True, space_before=1, space_after=2)
    add_p("(An Autonomous Institute Affiliated to Savitribai Phule Pune University, Pune)", size=11, bold=False, space_before=1, space_after=4)
    add_p("2025–2026", size=12, bold=True, space_before=2, space_after=0)
    
    # ---------------------------------------------------------
    # SECTION 2: CERTIFICATE PAGE
    # ---------------------------------------------------------
    sec2 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec2.top_margin = Pt(40)
    sec2.bottom_margin = Pt(14)
    sec2.left_margin = Pt(54)
    sec2.right_margin = Pt(54)
    
    # Logo on Certificate
    p_cert_logo = doc.add_paragraph()
    p_cert_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_logo.paragraph_format.space_before = Pt(0)
    p_cert_logo.paragraph_format.space_after = Pt(4)
    run_cert_logo = p_cert_logo.add_run()
    run_cert_logo.add_picture("college_logo.jpeg", width=Inches(1.05))
    
    add_p("CERTIFICATE", size=14, bold=True, space_before=4, space_after=10)
    
    add_p("This is to certify that the project report entitled", size=12, bold=False, space_before=4, space_after=4)
    add_p("“CAMPUS RESOURCE INTELLIGENCE & SHARING SYSTEM”", size=13, bold=True, space_before=2, space_after=8)
    
    add_p("Submitted by", size=12, bold=False, space_before=2, space_after=4)
    add_p("[Student Name]                  RBT No: [RBT Number]", size=12, bold=True, space_before=2, space_after=8)
    
    cert_text = (
        "is a bonafide student of this institute and the work has been carried out by him/her "
        "during the Second Year of the Bachelor of Technology (Information Technology) under the "
        "supervision of [Guide Name]. This work is approved for the partial fulfillment of the "
        "requirements of Savitribai Phule Pune University, Pune."
    )
    add_p(cert_text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=4, space_after=18, line_spacing=1.2)
    
    # Signatures Table (Guide & HOD)
    tbl_sigs = doc.add_table(rows=3, cols=2)
    tbl_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sigs.autofit = False
    
    # Borderless styling
    for row in tbl_sigs.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:top w:val="none"/>\n'
                f'  <w:left w:val="none"/>\n'
                f'  <w:bottom w:val="none"/>\n'
                f'  <w:right w:val="none"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)
    
    tbl_sigs.columns[0].width = Inches(3.4)
    tbl_sigs.columns[1].width = Inches(3.4)
    
    def set_cell(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=0):
        p = cell.paragraphs[0]
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)

    set_cell(tbl_sigs.cell(0, 0), "([Guide Name])", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_sigs.cell(0, 1), "(Dr. N. M. Ranjan)", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    set_cell(tbl_sigs.cell(1, 0), "Guide", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_sigs.cell(1, 1), "Head,", bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    set_cell(tbl_sigs.cell(2, 0), "Department of Information Technology", bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_sigs.cell(2, 1), "Department of Information Technology", bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    # Director block
    p_dir_space = doc.add_paragraph()
    p_dir_space.paragraph_format.space_before = Pt(14)
    p_dir_space.paragraph_format.space_after = Pt(0)
    
    tbl_dir = doc.add_table(rows=4, cols=2)
    tbl_dir.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_dir.autofit = False
    tbl_dir.columns[0].width = Inches(3.0)
    tbl_dir.columns[1].width = Inches(3.8)
    
    for row in tbl_dir.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>\n'
                f'  <w:top w:val="none"/>\n'
                f'  <w:left w:val="none"/>\n'
                f'  <w:bottom w:val="none"/>\n'
                f'  <w:right w:val="none"/>\n'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    set_cell(tbl_dir.cell(0, 0), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_dir.cell(0, 1), "(Dr. A. M. Badadhe)", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    set_cell(tbl_dir.cell(1, 0), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_dir.cell(1, 1), "Director,", bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    set_cell(tbl_dir.cell(2, 0), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_dir.cell(2, 1), "JSPM’s Rajarshi Shahu College of Engineering Pune – 33", bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    set_cell(tbl_dir.cell(3, 0), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell(tbl_dir.cell(3, 1), "(An Autonomous Institute Affiliated to Savitribai Phule Pune University, Pune)", bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT)
    
    # Place & Date
    p_pd = doc.add_paragraph()
    p_pd.paragraph_format.space_before = Pt(14)
    p_pd.paragraph_format.space_after = Pt(0)
    run_pd1 = p_pd.add_run("Place : Pune                                                                              Date : [Date]")
    run_pd1.font.name = 'Times New Roman'
    run_pd1.font.size = Pt(11)
    
    # ---------------------------------------------------------
    # SECTION 3: ABSTRACT PAGE
    # ---------------------------------------------------------
    sec3 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec3.top_margin = Pt(72)
    sec3.bottom_margin = Pt(72)
    sec3.left_margin = Pt(72)
    sec3.right_margin = Pt(72)
    
    add_p("ABSTRACT", size=14, bold=True, space_before=12, space_after=18)
    
    abstract_text = (
        "Campus Resource Intelligence and Sharing System (CRISS) addresses the systemic fragmentation "
        "of academic hardware, lab equipment, textbooks, and study spaces across institutional departments. "
        "In higher education environments, decentralized resource tracking leads to artificial shortages, "
        "prohibitive student costs, and severe underutilization of high-value equipment. CRISS resolves this "
        "through an enterprise relational architecture that unifies institutional inventory workflows with a "
        "student-to-student peer lending exchange. Built upon a 3NF-compliant MySQL 8 database with Python "
        "Flask, the system provides real-time transactional checkout, reservation scheduling, policy-governed "
        "peer listings, and an objective unmet-demand allocation engine. Empirical validation demonstrates "
        "strict transaction isolation, zero concurrent reservation collisions, and actionable visibility "
        "into resource demand deficits."
    )
    
    p_abs = add_p(abstract_text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=6, space_after=12, line_spacing=1.35)
    
    # Save document
    doc.save(output_path)
    print(f"Saved {output_path} successfully!")
    
    # Verify word count of abstract
    words = abstract_text.split()
    print(f"Abstract word count: {len(words)} (Required strictly 100-120 words)")
    assert 100 <= len(words) <= 120, f"Abstract word count {len(words)} out of range!"

if __name__ == "__main__":
    create_cover_certificate_abstract()
